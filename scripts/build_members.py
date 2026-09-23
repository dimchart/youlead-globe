#!/usr/bin/env python3
"""Перетворює data/members.xlsx на members.js, який читає сайт.

Запуск локально:   python3 scripts/build_members.py
(потрібен пакет openpyxl:  pip install openpyxl)

Якщо в Excel не вказані широта/довгота, скрипт шукає місто через OpenStreetMap
Nominatim і запам'ятовує результат у data/geocache.json.

Поки в Excel щось не так (немає файлу, порожній список), існуючий members.js
НЕ перезаписується — сайт продовжує показувати попередні дані.
"""
import json
import re
import sys
import time
import urllib.parse
import urllib.request
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent.parent
XLSX = ROOT / "data" / "members.xlsx"
CACHE = ROOT / "data" / "geocache.json"
OUT = ROOT / "members.js"
SHEET = "Учасники"

# Назви колонок, які ми розпізнаємо (регістр і зайві пробіли не важливі)
COLUMNS = {
    "name": ["ім'я", "імя", "ім’я", "name", "піб"],
    "role": ["підпис (потік / роль)", "підпис", "роль", "потік", "role"],
    "city": ["місто", "city"],
    "country": ["країна", "country"],
    "telegram": ["telegram", "телеграм", "tg"],
    "lat": ["широта", "lat", "latitude"],
    "lng": ["довгота", "lng", "lon", "longitude"],
}

NOMINATIM = "https://nominatim.openstreetmap.org/search"
USER_AGENT = "youlead-globe-members-builder/1.0 (https://github.com/dimchart/youlead-globe)"


def fail(msg):
    print(f"ПОМИЛКА: {msg}", file=sys.stderr)
    print("members.js не змінено.", file=sys.stderr)
    sys.exit(1)


def norm(s):
    return re.sub(r"\s+", " ", str(s or "")).strip().lower()


def clean(v):
    return re.sub(r"\s+", " ", str(v)).strip() if v is not None else ""


def to_float(v):
    """Приймає число або текст із комою ('50,45')."""
    if v is None or v == "":
        return None
    try:
        return float(str(v).strip().replace(",", "."))
    except ValueError:
        return None


def clean_telegram(v):
    s = clean(v)
    s = re.sub(r"^(https?://)?(www\.)?(t\.me|telegram\.me)/", "", s, flags=re.I)
    s = s.lstrip("@").split("?")[0].strip("/ ")
    return s


def load_cache():
    if CACHE.exists():
        try:
            return json.loads(CACHE.read_text(encoding="utf-8"))
        except json.JSONDecodeError:
            print("Увага: geocache.json пошкоджений, починаю з порожнього кешу.")
    return {}


def geocode(city, country, cache, state):
    key = f"{norm(city)}|{norm(country)}"
    if key in cache:
        return cache[key]["lat"], cache[key]["lng"]

    # Nominatim просить не частіше одного запиту на секунду
    wait = 1.1 - (time.time() - state["last"])
    if wait > 0:
        time.sleep(wait)
    query = ", ".join(p for p in (city, country) if p)
    url = NOMINATIM + "?" + urllib.parse.urlencode(
        {"q": query, "format": "json", "limit": 1, "accept-language": "uk"}
    )
    req = urllib.request.Request(url, headers={"User-Agent": USER_AGENT})
    try:
        with urllib.request.urlopen(req, timeout=20) as resp:
            data = json.load(resp)
    except Exception as e:  # мережа, ліміти тощо — не валимо весь запуск
        state["last"] = time.time()
        print(f"  не вдалося знайти «{query}»: {e}")
        return None
    state["last"] = time.time()
    if not data:
        print(f"  місто «{query}» не знайдено")
        return None
    lat, lng = round(float(data[0]["lat"]), 4), round(float(data[0]["lon"]), 4)
    cache[key] = {"lat": lat, "lng": lng}
    print(f"  знайдено «{query}» → {lat}, {lng}")
    return lat, lng


def main():
    if not XLSX.exists():
        fail(f"не знайдено файл {XLSX.relative_to(ROOT)}")

    try:
        wb = load_workbook(XLSX, data_only=True, read_only=True)
    except Exception as e:
        fail(f"не вдалося відкрити Excel: {e}")
    ws = wb[SHEET] if SHEET in wb.sheetnames else wb[wb.sheetnames[0]]

    rows = ws.iter_rows(values_only=True)
    header = next(rows, None)
    if not header:
        fail("аркуш порожній")

    index = {}
    for field, names in COLUMNS.items():
        for i, h in enumerate(header):
            if norm(h) in names:
                index[field] = i
                break
    for required in ("name", "city"):
        if required not in index:
            fail(f"немає колонки «{COLUMNS[required][0]}» у першому рядку")

    def cell(row, field):
        i = index.get(field)
        return row[i] if i is not None and i < len(row) else None

    cache = load_cache()
    state = {"last": 0.0}
    members, skipped = [], 0

    for n, row in enumerate(rows, start=2):
        name = clean(cell(row, "name"))
        if not name:
            continue  # порожній рядок
        city = clean(cell(row, "city"))
        country = clean(cell(row, "country"))
        lat, lng = to_float(cell(row, "lat")), to_float(cell(row, "lng"))

        if lat is None or lng is None:
            if not city:
                print(f"Рядок {n} ({name}): немає ні міста, ні координат — пропущено")
                skipped += 1
                continue
            found = geocode(city, country, cache, state)
            if not found:
                print(f"Рядок {n} ({name}): координати не знайдено — пропущено")
                skipped += 1
                continue
            lat, lng = found

        if not (-90 <= lat <= 90 and -180 <= lng <= 180):
            print(f"Рядок {n} ({name}): координати поза межами — пропущено")
            skipped += 1
            continue

        telegram = clean_telegram(cell(row, "telegram"))
        if telegram and not re.fullmatch(r"[A-Za-z0-9_]{3,32}", telegram):
            print(f"Рядок {n} ({name}): дивний Telegram-нік «{telegram}» — перевір")

        m = {"name": name}
        role = clean(cell(row, "role"))
        if role:
            m["role"] = role
        m["city"] = city or country or "—"
        m["lat"] = round(lat, 4)
        m["lng"] = round(lng, 4)
        if telegram:
            m["telegram"] = telegram
        members.append(m)

    if not members:
        fail("у файлі не знайдено жодного учасника (захист від випадкового очищення сайту)")

    CACHE.write_text(json.dumps(cache, ensure_ascii=False, indent=2, sort_keys=True) + "\n", encoding="utf-8")

    body = json.dumps(members, ensure_ascii=False, indent=2)
    content = (
        "// АВТОМАТИЧНО ЗГЕНЕРОВАНО зі скрипта scripts/build_members.py.\n"
        "// Не редагуй цей файл вручну — змінюй data/members.xlsx.\n"
        f"window.MEMBERS = {body};\n"
    )
    old = OUT.read_text(encoding="utf-8") if OUT.exists() else ""
    if content == old:
        print(f"Без змін: {len(members)} учасників (пропущено рядків: {skipped}).")
    else:
        OUT.write_text(content, encoding="utf-8")
        print(f"Оновлено members.js: {len(members)} учасників (пропущено рядків: {skipped}).")


if __name__ == "__main__":
    main()
